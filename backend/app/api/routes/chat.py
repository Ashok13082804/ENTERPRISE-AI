"""Chat API Routes — Ollama + RAG"""
import time
from typing import List, Optional
from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from sqlalchemy.ext.asyncio import AsyncSession
from pydantic import BaseModel

from app.core.database import get_db
from app.core.security import get_current_user
from app.ai.ollama_client import ollama_client
from app.ai.rag_pipeline import rag_service
from app.ai.writing_service import ai_writing_service
from app.models.chat import ChatSession, ChatMessage

router = APIRouter()


class ChatRequest(BaseModel):
    message: str
    session_id: Optional[int] = None
    model: str = "llama3"
    mode: str = "chat"  # chat, rag, code, sql, document
    rag_collection: Optional[str] = None
    stream: bool = False
    temperature: float = 0.7


class SessionCreate(BaseModel):
    title: Optional[str] = None
    model: str = "llama3"
    mode: str = "chat"
    rag_collection: Optional[str] = None


@router.get("/sessions")
async def list_sessions(
    db: AsyncSession = Depends(get_db),
    current_user=Depends(get_current_user),
):
    """List all chat sessions for current user."""
    from sqlalchemy import select, desc
    result = await db.execute(
        select(ChatSession)
        .where(ChatSession.user_id == current_user.id, ChatSession.is_archived == False)
        .order_by(desc(ChatSession.updated_at))
        .limit(50)
    )
    sessions = result.scalars().all()
    return [
        {
            "id": s.id,
            "title": s.title or "New Chat",
            "model": s.model,
            "mode": s.mode,
            "is_rag_enabled": s.is_rag_enabled,
            "created_at": s.created_at.isoformat(),
            "updated_at": s.updated_at.isoformat() if s.updated_at else None,
        }
        for s in sessions
    ]


@router.post("/sessions")
async def create_session(
    body: SessionCreate,
    db: AsyncSession = Depends(get_db),
    current_user=Depends(get_current_user),
):
    """Create a new chat session."""
    session = ChatSession(
        user_id=current_user.id,
        title=body.title,
        model=body.model,
        mode=body.mode,
        is_rag_enabled=body.rag_collection is not None,
        rag_collection=body.rag_collection,
    )
    db.add(session)
    await db.flush()
    return {"id": session.id, "title": session.title, "model": session.model}


@router.delete("/sessions/clear-all")
async def clear_all_history(
    db: AsyncSession = Depends(get_db),
    current_user=Depends(get_current_user),
):
    """Clear all chat sessions and messages for the current user."""
    from sqlalchemy import delete, select
    sessions_result = await db.execute(
        select(ChatSession.id).where(ChatSession.user_id == current_user.id)
    )
    session_ids = [row for row, in sessions_result.all()]
    if session_ids:
        await db.execute(
            delete(ChatMessage).where(ChatMessage.session_id.in_(session_ids))
        )
        await db.execute(
            delete(ChatSession).where(ChatSession.id.in_(session_ids))
        )
    await db.commit()
    return {"message": "All chat history cleared successfully"}


@router.get("/sessions/{session_id}/messages")
async def get_messages(
    session_id: int,
    db: AsyncSession = Depends(get_db),
    current_user=Depends(get_current_user),
):
    """Get all messages in a session."""
    from sqlalchemy import select
    result = await db.execute(
        select(ChatMessage)
        .where(ChatMessage.session_id == session_id)
        .order_by(ChatMessage.created_at)
    )
    messages = result.scalars().all()
    return [
        {
            "id": m.id,
            "role": m.role,
            "content": m.content,
            "model": m.model,
            "sources": m.sources,
            "processing_time_ms": m.processing_time_ms,
            "created_at": m.created_at.isoformat(),
        }
        for m in messages
    ]


@router.post("/send")
async def send_message(
    body: ChatRequest,
    db: AsyncSession = Depends(get_db),
    current_user=Depends(get_current_user),
):
    """Send a message and get AI response."""
    start_time = time.time()
    
    # Get or create session
    session_id = body.session_id
    if not session_id:
        session = ChatSession(
            user_id=current_user.id,
            title=body.message[:50],
            model=body.model,
            mode=body.mode,
            is_rag_enabled=body.rag_collection is not None,
            rag_collection=body.rag_collection,
        )
        db.add(session)
        await db.flush()
        session_id = session.id
    
    # Save user message
    user_msg = ChatMessage(
        session_id=session_id,
        role="user",
        content=body.message,
        model=body.model,
    )
    db.add(user_msg)
    
    # Get chat history
    from sqlalchemy import select
    hist_result = await db.execute(
        select(ChatMessage)
        .where(ChatMessage.session_id == session_id)
        .order_by(ChatMessage.created_at)
        .limit(20)
    )
    history = hist_result.scalars().all()
    messages = [{"role": m.role, "content": m.content} for m in history]
    messages.append({"role": "user", "content": body.message})
    
    sources = []
    
    if body.mode == "rag" and body.rag_collection:
        # RAG mode
        result = await rag_service.query(
            question=body.message,
            collection_name=body.rag_collection,
            model=body.model,
            chat_history=messages[:-1],
        )
        response_text = result["answer"]
        sources = result.get("sources", [])
    else:
        # Direct chat mode
        if body.mode == "code":
            system = "You are an expert software engineer. Write clean, production-ready code with comments."
        elif body.mode == "sql":
            system = "You are a senior database engineer. Write optimized SQL queries with explanations."
        else:
            system = "You are an intelligent enterprise AI assistant. Be professional, accurate, and helpful."
        
        full_messages = [{"role": "system", "content": system}] + messages
        response_text = await ollama_client.chat(
            messages=full_messages,
            model=body.model,
            temperature=body.temperature,
        )
    
    elapsed_ms = int((time.time() - start_time) * 1000)
    
    # Save assistant message
    ai_msg = ChatMessage(
        session_id=session_id,
        role="assistant",
        content=response_text,
        model=body.model,
        sources=sources,
        processing_time_ms=elapsed_ms,
    )
    db.add(ai_msg)
    await db.flush()
    
    return {
        "session_id": session_id,
        "message_id": ai_msg.id,
        "response": response_text,
        "sources": sources,
        "model": body.model,
        "processing_time_ms": elapsed_ms,
    }


@router.get("/models")
async def list_models():
    """List available Ollama models."""
    models = await ollama_client.list_models()
    is_available = await ollama_client.is_available()
    return {
        "available": is_available,
        "models": models,
        "ollama_url": ollama_client.base_url,
    }


@router.delete("/sessions/{session_id}")
async def delete_session(
    session_id: int,
    db: AsyncSession = Depends(get_db),
    current_user=Depends(get_current_user),
):
    from sqlalchemy import select
    result = await db.execute(
        select(ChatSession).where(
            ChatSession.id == session_id,
            ChatSession.user_id == current_user.id
        )
    )
    session = result.scalar_one_or_none()
    if not session:
        raise HTTPException(404, "Session not found")
    session.is_archived = True
    return {"message": "Session archived"}


@router.get("/sessions/{session_id}/export")
async def export_chat_session(
    session_id: int,
    format: str = "pdf",
    db: AsyncSession = Depends(get_db),
    current_user=Depends(get_current_user),
):
    """Export all chat messages in a session to PDF, Word, Markdown, Text, or JSON."""
    from sqlalchemy import select
    from fastapi.responses import Response
    import io
    
    # Verify session ownership
    result = await db.execute(
        select(ChatSession).where(
            ChatSession.id == session_id,
            ChatSession.user_id == current_user.id
        )
    )
    session = result.scalar_one_or_none()
    if not session:
        raise HTTPException(404, "Session not found")
        
    # Get messages
    msg_result = await db.execute(
        select(ChatMessage)
        .where(ChatMessage.session_id == session_id)
        .order_by(ChatMessage.created_at)
    )
    messages = msg_result.scalars().all()
    
    title = session.title or f"Chat_Session_{session_id}"
    # Replace spaces for filename safety
    safe_title = title.replace(" ", "_")
    
    if format == "json":
        import json
        data = [
            {"role": m.role, "content": m.content, "timestamp": m.created_at.isoformat()}
            for m in messages
        ]
        return Response(
            content=json.dumps(data, indent=2),
            media_type="application/json",
            headers={"Content-Disposition": f"attachment; filename={safe_title}.json"}
        )
        
    elif format == "txt":
        lines = [f"=== {title} ===", f"Exported: {time.strftime('%Y-%m-%d %H:%M:%S')}\n"]
        for m in messages:
            role_name = "User" if m.role == "user" else "Assistant"
            lines.append(f"[{role_name}]: {m.content}\n")
        content = "\n".join(lines)
        return Response(
            content=content,
            media_type="text/plain",
            headers={"Content-Disposition": f"attachment; filename={safe_title}.txt"}
        )
        
    elif format in ["markdown", "md"]:
        lines = [f"# {title}", f"*Exported: {time.strftime('%Y-%m-%d %H:%M:%S')}*\n"]
        for m in messages:
            role_name = "**User**" if m.role == "user" else "**Assistant**"
            lines.append(f"### {role_name}\n\n{m.content}\n")
        content = "\n".join(lines)
        return Response(
            content=content,
            media_type="text/markdown",
            headers={"Content-Disposition": f"attachment; filename={safe_title}.md"}
        )
        
    elif format == "docx":
        try:
            from docx import Document
            doc = Document()
            doc.add_heading(title, level=0)
            doc.add_paragraph(f"Exported: {time.strftime('%Y-%m-%d %H:%M:%S')}")
            doc.add_paragraph("")
            
            for m in messages:
                role_name = "User" if m.role == "user" else "Assistant"
                p = doc.add_paragraph()
                r = p.add_run(f"{role_name}: ")
                r.bold = True
                p.add_run(m.content)
                doc.add_paragraph("")
                
            buffer = io.BytesIO()
            doc.save(buffer)
            buffer.seek(0)
            return Response(
                content=buffer.getvalue(),
                media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                headers={"Content-Disposition": f"attachment; filename={safe_title}.docx"}
            )
        except Exception as e:
            raise HTTPException(500, f"DOCX generation error: {str(e)}")
            
    elif format == "pdf":
        try:
            from reportlab.lib.pagesizes import letter
            from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
            from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
            from reportlab.lib import colors
            
            buffer = io.BytesIO()
            pdf = SimpleDocTemplate(
                buffer,
                pagesize=letter,
                rightMargin=54,
                leftMargin=54,
                topMargin=54,
                bottomMargin=54
            )
            
            styles = getSampleStyleSheet()
            title_style = ParagraphStyle(
                'TitleStyle',
                parent=styles['Heading1'],
                fontSize=18,
                textColor=colors.HexColor('#4F46E5'),
                spaceAfter=15
            )
            meta_style = ParagraphStyle(
                'MetaStyle',
                parent=styles['Normal'],
                fontSize=9,
                textColor=colors.HexColor('#6B7280'),
                spaceAfter=25
            )
            user_style = ParagraphStyle(
                'UserStyle',
                parent=styles['Normal'],
                fontSize=11,
                textColor=colors.HexColor('#065F46'),
                backColor=colors.HexColor('#ECFDF5'),
                spaceAfter=15,
            )
            assistant_style = ParagraphStyle(
                'AssistantStyle',
                parent=styles['Normal'],
                fontSize=11,
                textColor=colors.HexColor('#1E1B4B'),
                backColor=colors.HexColor('#F3F4F6'),
                spaceAfter=15,
            )
            
            story = []
            story.append(Paragraph(title, title_style))
            story.append(Paragraph(f"Exported: {time.strftime('%Y-%m-%d %H:%M:%S')}", meta_style))
            
            for m in messages:
                role_name = "User" if m.role == "user" else "Assistant"
                style = user_style if m.role == "user" else assistant_style
                text = m.content.replace("\n", "<br/>").replace("<", "&lt;").replace(">", "&gt;")
                formatted_text = f"<b>{role_name}:</b><br/>{text}"
                story.append(Paragraph(formatted_text, style))
                story.append(Spacer(1, 10))
                
            pdf.build(story)
            buffer.seek(0)
            return Response(
                content=buffer.getvalue(),
                media_type="application/pdf",
                headers={"Content-Disposition": f"attachment; filename={safe_title}.pdf"}
            )
        except Exception as e:
            raise HTTPException(500, f"PDF generation error: {str(e)}")
            
    else:
        raise HTTPException(400, f"Unsupported format: {format}")


@router.post("/sessions/clear")
async def clear_all_sessions(
    db: AsyncSession = Depends(get_db),
    current_user=Depends(get_current_user),
):
    """Archive all chat sessions for the current user."""
    from sqlalchemy import update
    await db.execute(
        update(ChatSession)
        .where(ChatSession.user_id == current_user.id)
        .values(is_archived=True)
    )
    return {"message": "All chat sessions archived"}
