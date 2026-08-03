"""
RAG Pipeline — Complete Retrieval-Augmented Generation System
Supports: PDF, DOCX, XLSX, CSV, PPTX, TXT, MD
Vector Store: ChromaDB (local)
Embeddings: Ollama nomic-embed-text (local)
"""
import hashlib
import json
import re
from pathlib import Path
from typing import List, Dict, Any, Optional, Tuple

from loguru import logger

from app.core.config import settings
from app.ai.ollama_client import ollama_client


# ─── Document Loaders ─────────────────────────────────────────────────────────
def load_pdf(file_path: str) -> str:
    """Extract text from PDF."""
    try:
        from pypdf import PdfReader
        reader = PdfReader(file_path)
        texts = []
        for page in reader.pages:
            t = page.extract_text()
            if t:
                texts.append(t.strip())
        return "\n\n".join(texts)
    except Exception as e:
        logger.error(f"PDF load error: {e}")
        return ""


def load_docx(file_path: str) -> str:
    """Extract text from Word document."""
    try:
        from docx import Document
        doc = Document(file_path)
        paras = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        return "\n\n".join(paras)
    except Exception as e:
        logger.error(f"DOCX load error: {e}")
        return ""


def load_xlsx(file_path: str) -> str:
    """Extract text from Excel."""
    try:
        import pandas as pd
        dfs = pd.read_excel(file_path, sheet_name=None)
        texts = []
        for sheet_name, df in dfs.items():
            texts.append(f"Sheet: {sheet_name}")
            texts.append(df.to_string(index=False))
        return "\n\n".join(texts)
    except Exception as e:
        logger.error(f"XLSX load error: {e}")
        return ""


def load_csv(file_path: str) -> str:
    """Extract text from CSV."""
    try:
        import pandas as pd
        df = pd.read_csv(file_path)
        return df.to_string(index=False)
    except Exception as e:
        logger.error(f"CSV load error: {e}")
        return ""


def load_pptx(file_path: str) -> str:
    """Extract text from PowerPoint."""
    try:
        from pptx import Presentation
        prs = Presentation(file_path)
        texts = []
        for i, slide in enumerate(prs.slides):
            texts.append(f"Slide {i+1}:")
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    texts.append(shape.text.strip())
        return "\n".join(texts)
    except Exception as e:
        logger.error(f"PPTX load error: {e}")
        return ""


def load_text(file_path: str) -> str:
    """Load plain text or markdown."""
    try:
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            return f.read()
    except Exception as e:
        logger.error(f"Text load error: {e}")
        return ""


def extract_text(file_path: str, file_type: str) -> str:
    """Route to appropriate loader based on file type."""
    # Resolve relative paths to absolute to prevent cwd issues
    abs_path = str(Path(file_path).resolve())
    
    loaders = {
        "pdf": load_pdf,
        "docx": load_docx,
        "doc": load_docx,
        "xlsx": load_xlsx,
        "xls": load_xlsx,
        "csv": load_csv,
        "pptx": load_pptx,
        "ppt": load_pptx,
        "txt": load_text,
        "md": load_text,
        "markdown": load_text,
    }
    ext = file_type.lower().strip(".")
    loader = loaders.get(ext, load_text)
    return loader(abs_path)


async def query_module_rag(
    question: str,
    system_prompt: str,
    collection_name: str,
    model: str = None,
    n_results: int = 3,
) -> str:
    """Helper to query a specific ChromaDB collection for context and feed it to Ollama."""
    try:
        retrieved = await vector_store.search(
            collection_name=collection_name,
            query=question,
            n_results=n_results,
        )
        context_parts = []
        for i, chunk in enumerate(retrieved):
            context_parts.append(f"[Context Source {i+1} from {chunk['metadata'].get('document_title', 'doc')}]: {chunk['text']}")
        context = "\n\n".join(context_parts)
    except Exception as e:
        logger.error(f"RAG context retrieval failed: {e}")
        context = ""
        
    messages = [
        {"role": "system", "content": f"{system_prompt}\n\nRetrieved context from knowledge base database:\n{context or 'No context available in the database.'}"},
        {"role": "user", "content": question}
    ]
    return await ollama_client.chat(messages, model=model)


# ─── Text Chunking ────────────────────────────────────────────────────────────
def chunk_text(
    text: str,
    chunk_size: int = 512,
    overlap: int = 64,
) -> List[Dict[str, Any]]:
    """Split text into overlapping chunks with metadata."""
    # Clean text
    text = re.sub(r'\n{3,}', '\n\n', text)
    text = re.sub(r' {2,}', ' ', text)
    
    # Sentence-aware splitting
    sentences = re.split(r'(?<=[.!?])\s+', text)
    
    chunks = []
    current_chunk = []
    current_size = 0
    chunk_index = 0
    
    for sentence in sentences:
        words = sentence.split()
        word_count = len(words)
        
        if current_size + word_count > chunk_size and current_chunk:
            chunk_text_str = " ".join(current_chunk)
            chunks.append({
                "index": chunk_index,
                "text": chunk_text_str,
                "word_count": current_size,
                "char_count": len(chunk_text_str),
            })
            chunk_index += 1
            
            # Overlap: keep last N words
            overlap_words = current_chunk[-overlap:] if overlap > 0 else []
            current_chunk = overlap_words + words
            current_size = len(current_chunk)
        else:
            current_chunk.extend(words)
            current_size += word_count
    
    # Last chunk
    if current_chunk:
        chunk_text_str = " ".join(current_chunk)
        chunks.append({
            "index": chunk_index,
            "text": chunk_text_str,
            "word_count": current_size,
            "char_count": len(chunk_text_str),
        })
    
    return chunks


# ─── Vector Store (ChromaDB) ─────────────────────────────────────────────────
class VectorStore:
    """Local ChromaDB vector store manager."""

    def __init__(self, persist_path: str = None):
        self.persist_path = persist_path or settings.CHROMA_DB_PATH
        self._client = None

    def get_client(self):
        if self._client is None:
            try:
                import chromadb
                self._client = chromadb.PersistentClient(path=self.persist_path)
            except ImportError:
                logger.error("ChromaDB not installed. Run: pip install chromadb")
                raise
        return self._client

    def get_or_create_collection(self, name: str):
        client = self.get_client()
        return client.get_or_create_collection(
            name=name,
            metadata={"hnsw:space": "cosine"},
        )

    def list_collections(self) -> List[str]:
        try:
            client = self.get_client()
            return [c.name for c in client.list_collections()]
        except Exception:
            return []

    def delete_collection(self, name: str) -> bool:
        try:
            client = self.get_client()
            client.delete_collection(name)
            return True
        except Exception:
            return False

    async def add_chunks(
        self,
        collection_name: str,
        chunks: List[Dict[str, Any]],
        document_id: int,
        document_title: str,
    ) -> int:
        """Add text chunks with embeddings to ChromaDB."""
        collection = self.get_or_create_collection(collection_name)
        added = 0
        
        for chunk in chunks:
            try:
                chunk_id = f"doc_{document_id}_chunk_{chunk['index']}"
                text = chunk["text"]
                
                # Get embedding from Ollama
                embedding = await ollama_client.generate_embedding(text)
                
                if embedding:
                    collection.add(
                        ids=[chunk_id],
                        embeddings=[embedding],
                        documents=[text],
                        metadatas=[{
                            "document_id": document_id,
                            "document_title": document_title,
                            "chunk_index": chunk["index"],
                            "word_count": chunk.get("word_count", 0),
                        }],
                    )
                    added += 1
                else:
                    # Fallback: add without embedding (ChromaDB will use its own)
                    collection.add(
                        ids=[chunk_id],
                        documents=[text],
                        metadatas=[{
                            "document_id": document_id,
                            "document_title": document_title,
                            "chunk_index": chunk["index"],
                        }],
                    )
                    added += 1
            except Exception as e:
                logger.error(f"Error adding chunk {chunk['index']}: {e}")
        
        return added

    async def search(
        self,
        collection_name: str,
        query: str,
        n_results: int = 5,
    ) -> List[Dict[str, Any]]:
        """Semantic search in vector store."""
        try:
            collection = self.get_or_create_collection(collection_name)
            count = collection.count()
            if count == 0:
                return []
            
            # Get query embedding
            query_embedding = await ollama_client.generate_embedding(query)
            
            if query_embedding:
                results = collection.query(
                    query_embeddings=[query_embedding],
                    n_results=min(n_results, count),
                    include=["documents", "metadatas", "distances"],
                )
            else:
                # Fallback text search
                results = collection.query(
                    query_texts=[query],
                    n_results=min(n_results, count),
                    include=["documents", "metadatas", "distances"],
                )
            
            output = []
            if results and results["documents"]:
                for doc, meta, dist in zip(
                    results["documents"][0],
                    results["metadatas"][0],
                    results["distances"][0],
                ):
                    output.append({
                        "text": doc,
                        "metadata": meta,
                        "score": 1 - float(dist),  # cosine distance -> similarity
                    })
            
            return output
        except Exception as e:
            logger.error(f"Vector search error: {e}")
            return []


# ─── RAG Service ──────────────────────────────────────────────────────────────
class RAGService:
    """Full RAG pipeline: ingest → index → retrieve → generate."""

    def __init__(self):
        self.vector_store = VectorStore()

    async def ingest_document(
        self,
        file_path: str,
        file_type: str,
        document_id: int,
        document_title: str,
        collection_name: str = "default",
    ) -> Dict[str, Any]:
        """Process and index a document into the vector store."""
        logger.info(f"Ingesting document: {document_title} (type={file_type})")
        
        # Extract text
        text = extract_text(file_path, file_type)
        if not text:
            return {"success": False, "error": "No text extracted", "chunks": 0}
        
        # Chunk
        chunks = chunk_text(text, chunk_size=512, overlap=64)
        logger.info(f"Created {len(chunks)} chunks from {len(text)} characters")
        
        # Add to vector store
        added = await self.vector_store.add_chunks(
            collection_name=collection_name,
            chunks=chunks,
            document_id=document_id,
            document_title=document_title,
        )
        
        # Generate document hash
        doc_hash = hashlib.sha256(text.encode()).hexdigest()
        
        return {
            "success": True,
            "chunks": added,
            "total_chars": len(text),
            "document_hash": doc_hash,
        }

    async def query(
        self,
        question: str,
        collection_name: str = "default",
        model: str = None,
        n_results: int = 5,
        chat_history: List[Dict] = None,
    ) -> Dict[str, Any]:
        """RAG query: retrieve relevant chunks and generate answer."""
        model = model or settings.DEFAULT_MODEL
        
        # Retrieve
        retrieved = await self.vector_store.search(
            collection_name=collection_name,
            query=question,
            n_results=n_results,
        )
        
        # Build context
        context_parts = []
        sources = []
        for i, chunk in enumerate(retrieved):
            context_parts.append(f"[Source {i+1}] {chunk['text']}")
            sources.append({
                "index": i + 1,
                "text": chunk["text"][:200] + "...",
                "metadata": chunk["metadata"],
                "score": round(chunk["score"], 3),
            })
        
        context = "\n\n".join(context_parts)
        
        # Build messages
        messages = [
            {
                "role": "system",
                "content": (
                    "You are an intelligent enterprise AI assistant. "
                    "Answer questions using ONLY the provided context. "
                    "If the context doesn't contain the answer, say so clearly. "
                    "Always cite the source numbers [Source N] in your answer.\n\n"
                    f"Context:\n{context}"
                ),
            }
        ]
        
        # Add chat history
        if chat_history:
            messages.extend(chat_history[-6:])  # Last 3 exchanges
        
        messages.append({"role": "user", "content": question})
        
        # Generate
        answer = await ollama_client.chat(messages=messages, model=model)
        
        return {
            "answer": answer,
            "sources": sources,
            "context_chunks": len(retrieved),
            "model": model,
        }


# Singleton instances
vector_store = VectorStore()
rag_service = RAGService()
