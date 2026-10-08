import os
import json
import csv
from pathlib import Path
from typing import Dict, Any
from app.core.config import EXPORT_DIR

try:
    from reportlab.lib.pagesizes import letter
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib import colors
except ImportError:
    pass

try:
    import docx
    from docx.shared import Inches, Pt, RGBColor
except ImportError:
    pass

class ReportExporter:
    @staticmethod
    def export_all(report_data: Dict[str, Any], base_filename: str) -> Dict[str, str]:
        """
        Exports the report into PDF, DOCX, HTML, JSON, CSV, and TXT formats.
        Returns a dictionary of download links/paths.
        """
        stem = Path(base_filename).stem.replace(" ", "_")
        pdf_file = EXPORT_DIR / f"{stem}_report.pdf"
        docx_file = EXPORT_DIR / f"{stem}_report.docx"
        html_file = EXPORT_DIR / f"{stem}_report.html"
        json_file = EXPORT_DIR / f"{stem}_report.json"
        csv_file = EXPORT_DIR / f"{stem}_modules.csv"
        txt_file = EXPORT_DIR / f"{stem}_report.txt"

        ReportExporter.export_json(report_data, str(json_file))
        ReportExporter.export_txt(report_data, str(txt_file))
        ReportExporter.export_csv(report_data, str(csv_file))
        ReportExporter.export_html(report_data, str(html_file))
        
        try:
            ReportExporter.export_docx(report_data, str(docx_file))
        except Exception as e:
            print(f"DOCX export error: {e}")

        try:
            ReportExporter.export_pdf(report_data, str(pdf_file))
        except Exception as e:
            print(f"PDF export error: {e}")

        return {
            "pdf": f"/api/exports/{pdf_file.name}",
            "docx": f"/api/exports/{docx_file.name}",
            "html": f"/api/exports/{html_file.name}",
            "json": f"/api/exports/{json_file.name}",
            "csv": f"/api/exports/{csv_file.name}",
            "txt": f"/api/exports/{txt_file.name}"
        }

    @staticmethod
    def export_json(report_data: Dict[str, Any], out_path: str):
        with open(out_path, "w", encoding="utf-8") as f:
            json.dump(report_data, f, indent=2)

    @staticmethod
    def export_txt(report_data: Dict[str, Any], out_path: str):
        lines = [
            "=" * 70,
            f"          {report_data.get('title', 'AI DOCUMENT ANALYSIS REPORT')}",
            "=" * 70,
            f"Generated: {report_data.get('generated_at', '')}",
            ""
        ]
        meta = report_data.get("metadata", {})
        lines.append("METADATA:")
        for k, v in meta.items():
            lines.append(f"  {k}: {v}")
        lines.append("-" * 70)

        for sec in report_data.get("sections", []):
            lines.append(f"\n[{sec.get('title')}]")
            cnt = sec.get("content")
            if isinstance(cnt, dict):
                for k, v in cnt.items():
                    lines.append(f"  * {k}: {v}")
            elif isinstance(cnt, list):
                for item in cnt:
                    if isinstance(item, dict):
                        lines.append(f"  * {item.get('id', '')} - {item.get('name', '')} [{item.get('status', '')}]: {item.get('summary', '')}")
                    else:
                        lines.append(f"  * {item}")
            else:
                lines.append(f"  {cnt}")

        with open(out_path, "w", encoding="utf-8") as f:
            f.write("\n".join(lines))

    @staticmethod
    def export_csv(report_data: Dict[str, Any], out_path: str):
        sections = report_data.get("sections", [])
        modules_sec = next((s for s in sections if "Module-wise" in s.get("title", "")), None)
        rows = modules_sec.get("content", []) if modules_sec else []

        with open(out_path, "w", newline="", encoding="utf-8") as f:
            writer = csv.writer(f)
            writer.writerow(["Module ID", "Module Name", "Category", "Confidence", "Status", "Summary"])
            for r in rows:
                if isinstance(r, dict):
                    writer.writerow([
                        r.get("id", ""),
                        r.get("name", ""),
                        r.get("category", ""),
                        r.get("confidence", ""),
                        r.get("status", ""),
                        r.get("summary", "")
                    ])

    @staticmethod
    def export_html(report_data: Dict[str, Any], out_path: str):
        meta = report_data.get("metadata", {})
        sections = report_data.get("sections", [])

        sections_html = []
        for s in sections:
            cnt = s.get("content")
            content_render = ""
            if isinstance(cnt, dict):
                content_render = "<ul class='meta-list'>" + "".join(f"<li><strong>{k}:</strong> {v}</li>" for k, v in cnt.items()) + "</ul>"
            elif isinstance(cnt, list):
                if cnt and isinstance(cnt[0], dict):
                    table_rows = "".join(f"<tr><td><code>{x.get('id')}</code></td><td>{x.get('name')}</td><td>{x.get('category')}</td><td><span class='badge'>{x.get('confidence')}</span></td><td>{x.get('summary')}</td></tr>" for x in cnt[:15])
                    content_render = f"<table class='data-table'><thead><tr><th>ID</th><th>Name</th><th>Category</th><th>Confidence</th><th>Summary</th></tr></thead><tbody>{table_rows}</tbody></table>"
                else:
                    content_render = "<ul class='item-list'>" + "".join(f"<li>{x}</li>" for x in cnt) + "</ul>"
            else:
                content_render = f"<p class='lead'>{cnt}</p>"

            sections_html.append(f"""
            <section class="report-section">
                <h2>{s.get('title')}</h2>
                <div class="section-body">{content_render}</div>
            </section>
            """)

        html_doc = f"""<!DOCTYPE html>
<html lang="en">
<head>
    <meta charset="UTF-8">
    <title>{report_data.get('title')}</title>
    <style>
        body {{ font-family: -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, sans-serif; background: #080b11; color: #f1f5f9; padding: 40px; margin: 0; }}
        .container {{ max-width: 1000px; margin: 0 auto; background: #0d1525; border: 1px solid #1e293b; border-radius: 12px; padding: 32px; box-shadow: 0 10px 30px rgba(0,0,0,0.5); }}
        h1 {{ color: #00f0ff; border-bottom: 2px solid #1e293b; padding-bottom: 16px; margin-top: 0; }}
        h2 {{ color: #38bdf8; font-size: 1.25rem; margin-top: 24px; border-left: 4px solid #00f0ff; padding-left: 12px; }}
        .meta-grid {{ display: grid; grid-template-columns: repeat(auto-fit, minmax(200px, 1fr)); gap: 16px; margin-bottom: 24px; }}
        .meta-card {{ background: #131d33; padding: 12px 16px; border-radius: 8px; border: 1px solid rgba(255,255,255,0.05); }}
        .meta-label {{ font-size: 0.75rem; color: #94a3b8; text-transform: uppercase; letter-spacing: 0.05em; }}
        .meta-val {{ font-size: 1.1rem; font-weight: bold; color: #f8fafc; margin-top: 4px; }}
        .data-table {{ width: 100%; border-collapse: collapse; margin-top: 12px; font-size: 0.88rem; }}
        .data-table th, .data-table td {{ border: 1px solid #1e293b; padding: 8px 12px; text-align: left; }}
        .data-table th {{ background: #131d33; color: #38bdf8; }}
        .badge {{ background: rgba(0, 240, 255, 0.15); color: #00f0ff; padding: 2px 8px; border-radius: 4px; font-size: 0.8rem; }}
        ul.item-list {{ padding-left: 20px; line-height: 1.6; }}
        p.lead {{ line-height: 1.6; color: #cbd5e1; font-size: 0.95rem; }}
    </style>
</head>
<body>
    <div class="container">
        <h1>{report_data.get('title')}</h1>
        <div class="meta-grid">
            <div class="meta-card"><div class="meta-label">File Name</div><div class="meta-val">{meta.get('file_name')}</div></div>
            <div class="meta-card"><div class="meta-label">Format</div><div class="meta-val">{meta.get('file_type')}</div></div>
            <div class="meta-card"><div class="meta-label">Document Type</div><div class="meta-val">{meta.get('document_type')}</div></div>
            <div class="meta-card"><div class="meta-label">Completed Modules</div><div class="meta-val">{meta.get('completed_modules_count')} / {meta.get('applicable_modules_count')}</div></div>
            <div class="meta-card"><div class="meta-label">Overall Confidence</div><div class="meta-val">{round(meta.get('overall_confidence', 0)*100, 1)}%</div></div>
        </div>
        {"".join(sections_html)}
    </div>
</body>
</html>"""
        with open(out_path, "w", encoding="utf-8") as f:
            f.write(html_doc)

    @staticmethod
    def export_docx(report_data: Dict[str, Any], out_path: str):
        doc = docx.Document()
        doc.add_heading(report_data.get("title", "AI Document Analysis Report"), 0)
        
        meta = report_data.get("metadata", {})
        doc.add_paragraph(f"Generated: {report_data.get('generated_at', '')}")
        doc.add_paragraph(f"Document Type: {meta.get('document_type', '')} | Format: {meta.get('file_type', '')} | Pages: {meta.get('pages', 1)}")

        for s in report_data.get("sections", []):
            doc.add_heading(s.get("title", ""), level=1)
            cnt = s.get("content")
            if isinstance(cnt, dict):
                for k, v in cnt.items():
                    doc.add_paragraph(f"{k}: {v}")
            elif isinstance(cnt, list):
                for item in cnt:
                    if isinstance(item, dict):
                        doc.add_paragraph(f"[{item.get('id')}] {item.get('name')} - {item.get('summary')}")
                    else:
                        doc.add_paragraph(str(item), style='List Bullet')
            else:
                doc.add_paragraph(str(cnt))

        doc.save(out_path)

    @staticmethod
    def export_pdf(report_data: Dict[str, Any], out_path: str):
        doc = SimpleDocTemplate(out_path, pagesize=letter)
        styles = getSampleStyleSheet()
        story = []

        title_style = ParagraphStyle(
            'TitleStyle',
            parent=styles['Heading1'],
            fontSize=18,
            textColor=colors.HexColor('#0284c7'),
            spaceAfter=14
        )
        body_style = ParagraphStyle(
            'BodyStyle',
            parent=styles['Normal'],
            fontSize=10,
            leading=14,
            textColor=colors.HexColor('#1e293b')
        )
        h2_style = ParagraphStyle(
            'Heading2Style',
            parent=styles['Heading2'],
            fontSize=12,
            textColor=colors.HexColor('#0369a1'),
            spaceBefore=10,
            spaceAfter=6
        )

        story.append(Paragraph(report_data.get("title", "AI Document Analysis Report"), title_style))
        meta = report_data.get("metadata", {})
        story.append(Paragraph(f"<b>Generated:</b> {report_data.get('generated_at')} | <b>Confidence:</b> {round(meta.get('overall_confidence', 0)*100, 1)}%", body_style))
        story.append(Spacer(1, 12))

        for s in report_data.get("sections", [])[:10]:
            story.append(Paragraph(s.get("title", ""), h2_style))
            cnt = s.get("content")
            if isinstance(cnt, dict):
                for k, v in cnt.items():
                    story.append(Paragraph(f"• <b>{k}:</b> {v}", body_style))
            elif isinstance(cnt, list):
                for item in cnt[:5]:
                    if isinstance(item, dict):
                        story.append(Paragraph(f"• [{item.get('id')}] {item.get('name')}: {item.get('summary')}", body_style))
                    else:
                        story.append(Paragraph(f"• {item}", body_style))
            else:
                story.append(Paragraph(str(cnt), body_style))
            story.append(Spacer(1, 8))

        doc.build(story)
